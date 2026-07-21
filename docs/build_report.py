"""Build the course report; generated document is verified by render_docx.py."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name("Tiny_Terminal_Agent_Report.docx")

def font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd"); node.set(qn("w:fill"), fill); props.append(node)

def set_cell_margins(cell):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None: mar = OxmlElement("w:tcMar"); tcPr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        item = OxmlElement(f"w:{side}"); item.set(qn("w:w"), "100" if side in ("top", "bottom") else "120"); item.set(qn("w:type"), "dxa"); mar.append(item)

def p(doc, text="", style=None, italic=False, bold=False):
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.line_spacing = 1.15
    r = para.add_run(text); font(r, 11, bold); r.italic = italic
    return para

def heading(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    para.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    para.paragraph_format.space_after = Pt(6)
    r = para.add_run(text); font(r, 16 if level == 1 else 13, True, "2E74B5")

def bullet(doc, text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(3); para.paragraph_format.line_spacing = 1.15
    font(para.add_run(text), 11)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""; shade(c, "E8EEF5"); set_cell_margins(c)
        r = c.paragraphs[0].add_run(h); font(r, 10, True, "1F4D78")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""; set_cell_margins(cells[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            font(cells[i].paragraphs[0].add_run(value), 9.5)
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths): cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def main():
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(.492)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    for s in ("Heading 1", "Heading 2"):
        doc.styles[s].font.name = "Calibri"
    header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("Tiny Terminal Agent | Mini-project Report"), 9, False, "666666")
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("Course mini-project | 2026"), 9, False, "666666")

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.paragraph_format.space_before = Pt(30); title.paragraph_format.space_after = Pt(10)
    font(title.add_run("Tiny Terminal Agent"), 18, True, "000000", name="Times New Roman")
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER; sub.paragraph_format.space_after = Pt(24)
    font(sub.add_run("Exploring Reliable Tool Use and Context Management for Small Language Models"), 14, False, "2E74B5")
    meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(meta.add_run("Mini-project report | Agent Systems"), 11, False, "555555")
    doc.add_page_break()

    heading(doc, "Abstract")
    p(doc, "Agent systems ask a language model to select tools, interpret observations, and continue until a task is complete. These behaviours are fragile for small language models because an unconstrained completion can contain prose, malformed arguments, or an incomplete memory of previous steps. This project implements Tiny Terminal Agent, a compact Python terminal agent that reduces the action space to two validated JSON actions: a shell-tool call or a final answer. The system combines schema validation with repair prompts, bounded rolling memory, truncated tool observations, a fixed step budget, and a local Ollama adapter. Local verification executed four unit tests covering protocol rejection, fenced-JSON normalisation, successful tool execution, and a destructive-command guard; all four passed. A real CPU evaluation with Qwen3 0.6B (751.63M parameters) compared a loose baseline with the optimised agent on three multi-step file tasks. Completion increased from 0/3 to 1/3, while valid-action rate remained 3/3 in both settings. The prototype demonstrates a practical boundary: strict interfaces can improve operational reliability for small models, but string-based command blocking is not a security sandbox and must be replaced by isolation and allowlists in production.")
    heading(doc, "1. Introduction")
    p(doc, "Language-model agents typically alternate between planning and tool use. The core challenge is not merely producing a plausible sentence; it is producing a machine-executable action at the right time, preserving relevant state, and stopping safely. Large proprietary models often provide dedicated tool-calling interfaces, whereas smaller locally deployed models may need a narrower and more explicit control protocol.")
    p(doc, "This mini-project investigates the following question: can a small-model-oriented terminal agent remain understandable and reproducible while improving the reliability of tool invocation? The answer explored here is architectural rather than model-specific. Tiny Terminal Agent constrains the output grammar, validates every action before execution, and keeps the interaction within an explicit resource budget.")
    heading(doc, "2. System Design")
    p(doc, "The one-sentence design claim is: Tiny Terminal Agent enables bounded terminal task execution by combining a two-action JSON protocol, validation-and-repair, and compacted context, with evidence from local unit and end-to-end tests. The scope is a classroom prototype running in a user-selected working directory; it does not claim adversarial containment or broad autonomous capability.")
    table(doc, ["Component", "Role", "Small-model-oriented optimisation"], [
        ("Protocol", "Accepts only tool or final actions", "Two templates reduce format ambiguity."),
        ("Validator", "Parses JSON and rejects unsupported keys", "Malformed output becomes a repair turn, not execution."),
        ("Memory", "Stores recent turns and compacted history", "Bounds prompt growth while preserving salient observations."),
        ("ShellTool", "Executes one shell command with a timeout", "Observations are truncated to 2,000 characters."),
        ("Adapters", "Connects DemoModel or local Ollama", "No third-party runtime dependency is required."),
    ], [1.2, 2.3, 3.0])
    heading(doc, "2.1 Action Protocol", 2)
    p(doc, "The model must output exactly one JSON object. A tool request has the form {\"type\":\"tool\",\"tool\":\"shell\",\"command\":\"...\"}; a completion has the form {\"type\":\"final\",\"answer\":\"...\"}. Any prose, invalid JSON, unknown key, unsupported tool, or empty command is rejected. The agent records a short repair instruction and requests another completion. This separates linguistic generation from execution authority.")
    heading(doc, "2.2 Context Management", 2)
    p(doc, "Each interaction is stored as a role-content pair. When the total character count exceeds 5,000, older entries are shortened into an eight-item summary while the most recent turns remain verbatim. Tool observations are capped at 2,000 characters. Together with the default six-step budget, these controls prevent an unbounded transcript from becoming the dominant input to a small model.")
    heading(doc, "2.3 Tool Interface and Safety Boundary", 2)
    p(doc, "The terminal tool launches a subprocess in the supplied workspace, applies a 20-second timeout, normalises output decoding as UTF-8 with replacement, and rejects a small set of plainly destructive command patterns. The guard is intentionally described as a usability safeguard rather than a sandbox: a language model can still construct unsafe commands not matched by a short denylist. A production executor should use operating-system isolation, read-only mounts where possible, capability-specific allowlists, and human approval for irreversible actions.")
    heading(doc, "3. Implementation")
    p(doc, "The implementation uses only the Python standard library for the agent path. protocol.py defines the schema; memory.py manages bounded state; llm.py provides the Ollama HTTP adapter and a deterministic demo model; tools.py encapsulates command execution; agent.py implements the reason-act-observe loop; and cli.py exposes a command-line interface. The actual evaluation used Qwen3 0.6B through Ollama. The adapter also accepts an explicit endpoint, which allowed a CPU-only service to be used after the default GPU runtime proved incompatible with the local CUDA driver.")
    heading(doc, "4. Evaluation")
    p(doc, "Evaluation combined executable correctness with a real small-model comparison. The Python unittest suite contained four cases: rejecting prose in place of JSON; accepting a fenced JSON object commonly emitted by small models; executing a valid echo command before accepting a final answer; and rejecting the explicit pattern rm -rf /. All four tests passed. For the model experiment, Qwen3 0.6B (reported by the local runtime as 751.63M parameters) was evaluated on three read-only Windows PowerShell tasks: locating a token in a file, counting .log files, and reading a configuration value. The baseline had only a loose COMMAND/FINAL prompt. The optimised condition added strict JSON actions, fenced-JSON normalisation, explicit PowerShell tool examples, a workspace-path constraint, result-driven retries, and bounded context.")
    table(doc, ["Check", "Observed result", "Interpretation"], [
        ("Unit tests", "4/4 PASS", "Protocol, fenced JSON, tool execution, and destructive-command guard."),
        ("Qwen3 0.6B baseline", "0/3 complete; 3/3 valid actions", "Loose command format did not yield a correct final task answer."),
        ("Qwen3 0.6B optimised", "1/3 complete; 3/3 valid actions", "The configuration-reading task returned mode=compact correctly."),
        ("Observed boundary", "2/3 still incomplete", "Token search exhausted the step budget; log counting returned an incorrect value."),
    ], [1.65, 1.05, 3.8])
    heading(doc, "4.1 Debugging Observation", 2)
    p(doc, "Two portability faults were discovered during real execution. First, Windows shell output caused a default GBK decoding failure; the subprocess wrapper was revised to decode explicitly as UTF-8 with replacement and to tolerate missing stdout or stderr values. Second, the default Ollama GPU runner failed with a CUDA 'device kernel image is invalid' error. A separate CPU-only Ollama service was launched on port 11435, and the adapter was extended with an explicit endpoint. This enabled the real Qwen3 0.6B evaluation. These incidents show that agent reliability includes the model runtime and tool-output layer, not only prompt wording.")
    heading(doc, "5. Discussion")
    p(doc, "The measured comparison supports a modest but useful conclusion. For a 751.63M-parameter model, explicit tool semantics and output normalisation changed task completion from 0/3 to 1/3 without fine-tuning. The gain is limited but meaningful because the tasks require grounding in real shell output rather than producing a plausible answer. The JSON parser, fenced-output normaliser, repair loop, bounded memory, and step budget turn several silent errors into visible control-flow events. The same choices also make the project easy to inspect: the agent has one tool, two actions, and a short execution trace.")
    p(doc, "The experiment also exposes limits. Valid-action rate alone was not sufficient: the baseline emitted the requested action format in all three cases but still failed every task, while the optimised agent failed token search by exhausting its step budget and returned an incorrect log count. The current safety mechanism is only a denylist. Future work should compare unconstrained prompting, JSON prompting, and grammar-constrained decoding on a larger task set; add a task-aware command planner; and use a containerised allowlisted executor. Useful metrics include valid-action rate, task completion rate, mean repair turns, context length, latency, and unsafe-command attempts.")
    heading(doc, "6. Responsible Use and AI Assistance")
    p(doc, "The system should be used only on a disposable or suitably isolated environment when model-generated commands are permitted. It should not be connected to private credentials, unrestricted network access, or valuable files without additional controls. AI tools were used to assist implementation planning, code drafting, and English technical writing. The documented implementation and its test results were then executed locally; no external benchmark result or model-performance figure is claimed without measurement.")
    heading(doc, "7. Conclusion")
    p(doc, "Tiny Terminal Agent is a working, compact Agent framework for exploring small-model tool use. It delivers a two-action protocol, fenced-output normalisation, validation repair, bounded context management, local Ollama integration, and reproducible tests. A real CPU evaluation of Qwen3 0.6B improved completion from 0/3 to 1/3 after prompt and tool-interface optimisation. The contribution is therefore a clear baseline for coursework and further experimentation, while its limits are explicit: 2/3 tasks remained incomplete, and operational security requires a real sandbox rather than pattern matching.")
    heading(doc, "Appendix A. Reproduction Commands")
    p(doc, "Source code repository: https://github.com/wangyusongs/tiny-terminal-agent", italic=True)
    p(doc, "python -m unittest discover -s tests -v", italic=True)
    p(doc, "python -m tinyagent.cli \"show the working directory\" --provider demo", italic=True)
    heading(doc, "Appendix B. Terminology Ledger")
    table(doc, ["Canonical term", "Definition", "Decision"], [
        ("Tiny Terminal Agent", "The implemented Python Agent framework", "Use this name throughout."),
        ("small language model", "A resource-constrained instruction-tuned model", "Avoid an unsupported exact parameter threshold."),
        ("tool action", "A validated request to invoke the shell tool", "Distinct from a final answer."),
        ("bounded context", "Recent turns plus compacted historical observations", "Implemented with a 5,000-character threshold."),
        ("DemoModel", "Deterministic offline adapter", "Used for reproducible verification, not performance measurement."),
    ], [1.6, 3.0, 1.9])
    doc.save(OUT)
    print(OUT)

if __name__ == "__main__": main()
